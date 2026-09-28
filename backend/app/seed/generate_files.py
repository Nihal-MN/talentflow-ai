"""Render the synthetic demo candidates to real PDF / DOCX / TXT files.

The seed command ingests these files through the SAME pipeline a recruiter
would use (text extraction → structured extraction → validation), so the demo
exercises the real path — nothing is inserted into the database directly.
"""

from __future__ import annotations

from pathlib import Path

from app.seed.demo_candidates import CANDIDATES
from app.seed.demo_jobs import DEMO_JOBS

#: PDFs use the built-in fonts, which only cover latin-1; map typographic chars.
_LATIN1_MAP = str.maketrans(
    {
        "\u2014": "-",
        "\u2013": "-",
        "\u2019": "'",
        "\u2018": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2022": "-",
        "\u00b7": "-",
        "\u2026": "...",
    }
)


def render_resume_text(candidate: dict) -> str:
    """Render one demo candidate to the canonical plain-text resume format."""
    lines: list[str] = [candidate["full_name"], candidate["headline"], candidate["location"]]
    contact = [candidate["email"], candidate["phone"], *candidate.get("links", [])]
    lines.append(" | ".join(contact))
    lines.append("")

    lines.extend(["Summary", candidate["summary"], ""])
    lines.append("Experience")
    for role in candidate["experience"]:
        lines.append(f"{role['title']} — {role['company']} ({role['location']})")
        lines.append(f"{role['start']} – {role['end']}")
        for bullet in role["bullets"]:
            lines.append(f"- {bullet}")
        lines.append("")
    lines.append("Education")
    lines.extend(candidate["education"])
    lines.append("")
    lines.append("Skills")
    lines.append(", ".join(candidate["skills"]))
    if candidate.get("certifications"):
        lines.append("")
        lines.append("Certifications")
        lines.extend(candidate["certifications"])
    lines.append("")
    return "\n".join(lines)


def _slug(name: str) -> str:
    return name.lower().replace(" ", "_").replace(".", "")


def write_txt(text: str, path: Path) -> None:
    path.write_text(text, encoding="utf-8")


def write_docx(text: str, path: Path) -> None:
    import docx

    document = docx.Document()
    for index, line in enumerate(text.split("\n")):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line)
        if index == 0 or line in ("Summary", "Experience", "Education", "Skills", "Certifications"):
            run.bold = True
    document.save(str(path))


def write_pdf(text: str, path: Path) -> None:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF(format="A4")
    pdf.set_margins(18, 15, 18)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10.5)
    for line in text.translate(_LATIN1_MAP).split("\n"):
        if not line.strip():
            pdf.ln(3.4)
            continue
        pdf.multi_cell(0, 5.2, text=line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.output(str(path))


def write_resume(candidate: dict, out_dir: Path) -> Path:
    """Write one candidate to its file format; returns the path."""
    text = render_resume_text(candidate)
    filename = f"{_slug(candidate['full_name'])}.resume.{candidate['format']}"
    path = out_dir / filename
    writers = {"pdf": write_pdf, "docx": write_docx, "txt": write_txt}
    writers[candidate["format"]](text, path)
    return path


def generate_demo_files(repo_root: Path, *, force: bool = False) -> dict[str, list[Path]]:
    """Generate ``examples/jobs/*.md`` and ``examples/resumes/*``.

    Existing files are kept unless ``force`` is set.
    """
    jobs_dir = repo_root / "examples" / "jobs"
    resumes_dir = repo_root / "examples" / "resumes"
    jobs_dir.mkdir(parents=True, exist_ok=True)
    resumes_dir.mkdir(parents=True, exist_ok=True)

    job_files: list[Path] = []
    for job in DEMO_JOBS.values():
        path = jobs_dir / job["filename"]
        if force or not path.exists():
            path.write_text(job["text"], encoding="utf-8")
        job_files.append(path)

    resume_files: list[Path] = []
    for candidate in CANDIDATES:
        filename = f"{_slug(candidate['full_name'])}.resume.{candidate['format']}"
        path = resumes_dir / filename
        if force or not path.exists():
            path = write_resume(candidate, resumes_dir)
        resume_files.append(path)

    return {"jobs": job_files, "resumes": resume_files}
